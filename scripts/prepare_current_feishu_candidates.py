from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SRC_ROOT = Path(__file__).resolve().parents[1] / "src"
if str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))

from mineru_pipeline.field_semantics import clean_variant_text, normalize_model_family
from mineru_pipeline.import_quality import build_quality_report, write_report_files


ROOT = Path("C:/Users/HP/Downloads/车源汇总纯净 (2)/车源汇总纯净")
SUBTOTAL_TOKENS = {"小计", "合计", "subtotal", "total"}
BASE_ONLY_PLACES = {"霍尔果斯", "湘潭", "盐城"}
NON_LOCATION_TEXT = {"可指定发运", "未知地点", "未知", "待定"}
BRAND_PATH_ALIASES = {
    "AVATR": {"Avatr", "AVATR", "阿维塔"},
    "BYD": {"BYD", "比亚迪"},
    "Changan": {"Changan", "长安"},
    "Deepal": {"Deepal", "深蓝"},
    "Dongfeng": {"Dongfeng", "东风"},
    "Farizon": {"Farizon", "Furuition", "远程", "福瑞通"},
    "GAC Aion": {"GAC", "GAC Aion", "广汽", "广汽埃安"},
    "Geely": {"Geely", "吉利"},
    "IM Motors": {"IM", "IM Motors", "SAIC", "智己"},
    "Leapmotor": {"Leapmotor", "零跑"},
    "Toyota": {"Toyota", "丰田"},
    "Wuling": {"Wuling", "五菱"},
    "Xiaomi": {"Xiaomi", "小米"},
}
VALID_BRANDS = set(BRAND_PATH_ALIASES)

CURRENT_FIELDS = [
    "model_id",
    "confidence",
    "brand",
    "supplier",
    "max_quantity",
    "exterior_color",
    "model",
    "variant",
    "supplier_price_cny",
    "cost_fca_usd",
    "trim_config_id",
    "cost_fob_usd",
    "notes",
    "order_wait_days",
    "manufacture_year",
    "min_quantity",
    "record_id",
    "interior_color",
    "location",
    "steering_setup",
    "market_region",
    "cost_exw_usd",
    "manufacture_month",
    "display_price_low",
    "stock_quantity",
    "vehicle_supply_base",
    "display_price_high",
]


def clean(value: Any) -> str | None:
    if value in (None, ""):
        return None
    text = str(value).strip()
    return text if text not in {"", "/", "-", "—"} else None


def split_place_fields(value: Any) -> tuple[str | None, str | None]:
    text = clean(value)
    if not text:
        return None, None
    without_trade = re.sub(r"^(EXW|FCA|FOB|CIF)\s*", "", text, flags=re.I).strip()
    if without_trade in NON_LOCATION_TEXT:
        return None, None
    if "基地" in without_trade or without_trade in BASE_ONLY_PLACES:
        return None, without_trade
    return without_trade, None


def number(value: Any) -> int | float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else float(value)
    match = re.search(r"-?\d+(?:\.\d+)?", str(value).replace(",", ""))
    if not match:
        return None
    parsed = float(match.group(0))
    return int(parsed) if parsed.is_integer() else parsed


def split_color_quantity(exterior: Any, stock: Any) -> tuple[int | float | None, str | None]:
    text = clean(exterior)
    qty = number(stock)
    if not text:
        return qty, None
    match = re.match(r"^(\d+)(.+)$", text)
    if match:
        return number(match.group(1)), clean(match.group(2))
    return qty, text


def normalize_wuling_model(model: str | None) -> str | None:
    return normalize_model_family(model, brand="Wuling")


def parse_year_month(value: Any) -> tuple[int | None, int | None]:
    if isinstance(value, datetime):
        return value.year, value.month
    text = clean(value)
    if not text:
        return None, None
    match = re.search(r"(20\d{2})[-年./](\d{1,2})", text)
    if match:
        return int(match.group(1)), int(match.group(2))
    match = re.search(r"(20\d{2})", text)
    if match:
        return int(match.group(1)), None
    return None, None


def normalize_market(value: Any) -> list[str] | None:
    text = clean(value)
    if not text:
        return None
    if "国际" in text or "出口" in text or "海外" in text:
        return ["国际版"]
    if "国内" in text or "中规" in text:
        return ["国内版"]
    return None


def split_model_trim(brand: str | None, raw_model: Any, raw_trim: Any = None) -> tuple[str | None, str | None]:
    model = clean(raw_model)
    trim = clean(raw_trim)
    if not model:
        return None, trim
    original = model
    if brand == "BYD" and trim and "海鸥" not in model and re.search(r"(Sealion|Seagull)", trim, re.I):
        return trim, model
    if brand == "Wuling":
        first_line = normalize_wuling_model(model.splitlines()[0].strip()) or model.splitlines()[0].strip()
        first_line = re.sub(r"\s*载重\s*\d+\s*kg.*$", "", first_line, flags=re.I).strip()
        if re.fullmatch(r"\d{3,5}", first_line):
            first_line = "Wuling"
        trim_parts = [part for part in (model if first_line != model else None, trim) if part]
        return first_line, " | ".join(trim_parts) if trim_parts else None
    if brand == "Geely":
        text = re.sub(r"^20\d{2}\s*Model\s*", "", model, flags=re.I)
        text = re.sub(r"^20\d{2}款\s*", "", text)
        geely_match = re.search(r"(Galaxy\s+[A-Za-z]+\s*\d+|Galaxy\s+[A-Z]\d+|Galaxy\s+XinYao\s+\d+|Cowboy|吉利牛仔|全新牛仔|Panda|Haoyue\s*L|BoYue\s+REV|Geely\s+ICON|A7|银河[^/\\n]+|博越REV|吉利ICON[^/\\n]*)", text, re.I)
        if geely_match:
            family = geely_match.group(1).strip(" /")
            if "牛仔" in family:
                family = "Cowboy"
            remainder = text.replace(geely_match.group(0), "", 1)
            remainder = re.sub(r"\bGeely\b", "", remainder, flags=re.I)
            remainder = re.sub(r"20\d{2}款?吉利银河\s*M?\s*\d+（[^）]*）", "", remainder)
            remainder = re.sub(r"吉利银河\s*M?\s*\d+", "", remainder)
            remainder = remainder.strip(" /|：:（）()")
            trim_parts = [part for part in (remainder, trim) if part]
            return family, " | ".join(trim_parts) if trim_parts else None
        if model in {"2025款", "2026款"}:
            return "全新牛仔", trim
    if brand == "Changan":
        text = re.sub(r"^20\d{2}款?", "", model).strip()
        for token in (
            "Qiyuan A06",
            "Qiyuan A07",
            "Qiyuan Q07",
            "L06",
            "Lumin",
            "S05",
            "S07",
            "S09",
            "UNI-Z PHEV",
            "UNI-Z",
            "CS75PLUS",
            "CS75 PRO",
            "长安X5PLUS",
            "X5PLUS",
            "第四代逸动",
        ):
            if token.lower() in text.lower():
                trim_parts = [part for part in (re.sub(re.escape(token), "", text, count=1, flags=re.I).strip(), trim) if part]
                return token, " | ".join(trim_parts) if trim_parts else None
    if brand == "AVATR":
        match = re.search(r"\b(Max|Ultra)\b", model, re.I)
        if match:
            family = match.group(1).title()
            remainder = model.replace(match.group(0), "", 1).strip(" -/：:（）()")
            trim_parts = [part for part in (remainder, trim) if part]
            return family, " | ".join(trim_parts) if trim_parts else None
    if brand == "IM Motors":
        match = re.search(r"\b(?:IM\s*)?(L6|LS6)\b", f"{model} {trim or ''}", re.I)
        if match:
            family = match.group(1).upper()
            remainder = re.sub(r"\b(?:IM\s*)?(L6|LS6)\b", "", f"{model} {trim or ''}", count=1, flags=re.I).strip(" -/：:（）()")
            return family, remainder or None
    if brand == "Xiaomi":
        match = re.search(r"\b(SU7|YU7)\b", model, re.I)
        if match:
            family = match.group(1).upper()
            remainder = model.replace(match.group(0), "", 1).strip(" -/：:（）()")
            trim_parts = [part for part in (remainder, trim) if part]
            return family, " | ".join(trim_parts) if trim_parts else None
    rules = [
        (r"\b(V6E|V7E|V8E|SV\d*|新V)\b", None),
        (r"(UNI-Z\s*PHEV|UNI-Z|CS75PLUS|第四代逸动|启源\s*Q05|Q05)", None),
        (r"(A7|全新牛仔|银河\s*E5|Galaxy\s*E5)", None),
        (r"(第二代秦PLUS|秦PLUS)", "秦PLUS"),
        (r"(驱逐舰05|Destroyer\s*05)", "Destroyer 05"),
        (r"(海狮05\s*EV|海狮05EV|Sealion\s*05\s*EV)", "Sealion 05 EV"),
        (r"(海狮06\s*EV|海狮06EV|Sealion\s*06\s*EV)", "Sealion 06 EV"),
        (r"(海狮06\s*DMI|海狮06DMI)", "Sealion 06 DMI"),
        (r"(海狮07\s*EV|海狮07EV|Sealion\s*7EV|Sealion\s*7)", "Sealion 7"),
        (r"(BYD\s+SHARK\s*6|SHARK\s*6)", "Shark 6"),
        (r"(BYD\s+TI7|TI7|钛7)", "Ti 7"),
        (r"(钛3|Tai\s*3|Ti\s*3)", "Ti 3"),
        (r"(ATTO3)", "ATTO3"),
        (r"(T03)", "T03"),
        (r"(Smart\s*#\d)", None),
        (r"(i60)", "i60"),
        (r"\b(L6|LS6)\b", None),
        (r"\b(X50|X70FL|X70PLUS|DASHING|T1|T2|G700)\b", None),
    ]
    for pattern, replacement in rules:
        match = re.search(pattern, model, re.I)
        if match:
            family = replacement or match.group(1).strip()
            remainder = model.replace(match.group(0), "", 1).strip(" -/：:（）()")
            trim_parts = [part for part in (remainder, trim) if part]
            return family, " | ".join(trim_parts) if trim_parts else None
    if brand == "Bestune" and "Pony" in model:
        trim_parts = [part for part in (model.replace("Bestune Pony", "", 1).replace("Pony", "", 1).strip(" -/：:（）()"), trim) if part]
        return "Xiaoma", " | ".join(trim_parts) if trim_parts else None
    if brand in {"Geely", "Wuling", "Changan", "Deepal", "Farizon", "AVATR", "Xiaomi", "IM Motors", "Toyota", "Foton", "Jetour", "Smart", "Bestune"}:
        return model, trim
    return original, trim


def normalize_variant_text(brand: str | None, model: str | None, variant: str | None) -> str | None:
    return clean_variant_text(variant, brand=brand, model=model)


def path_model_hint(source_file: str, brand: str | None) -> str | None:
    if not source_file or not brand:
        return None
    parts = re.split(r"[\\/]+", source_file)
    aliases = BRAND_PATH_ALIASES.get(brand, {brand})
    for index, part in enumerate(parts[:-1]):
        if part in aliases:
            candidate = clean(parts[index + 1])
            if candidate and not re.search(r"^(EXW|FCA|FOB|CIF)\b|价格表|原表|按工作簿", candidate, re.I):
                if brand == "Wuling":
                    return normalize_wuling_model(candidate)
                return candidate
    return None


def path_brand_hint(source_file: str) -> str | None:
    parts = re.split(r"[\\/]+", source_file or "")
    for part in parts:
        for brand, aliases in BRAND_PATH_ALIASES.items():
            if part in aliases:
                return brand
    return None


def evidence_text(row: dict[str, Any]) -> str:
    source_file = row.get("source_file")
    sheet = row.get("source_sheet")
    source_row = row.get("source_row")
    parts = [f"# Path: {source_file}", f"Sheet: {sheet}", f"row {source_row}"]
    if source_file and sheet and source_row:
        path = ROOT / str(source_file)
        try:
            wb = load_workbook(path, data_only=True)
            ws = wb[str(sheet)]
            header_values = [cell.value for cell in ws[1]]
            header = " | ".join(str(v) for v in header_values if v not in (None, ""))
            if header:
                parts.append(f"header: {header}")
            values = [cell.value for cell in ws[int(source_row)]]
            rendered = " | ".join(str(v) for v in values if v not in (None, ""))
            if rendered:
                parts.append(rendered)
            wb.close()
        except Exception:
            pass
    return "; ".join(str(part) for part in parts if part)


def source_values(row: dict[str, Any]) -> list[Any]:
    source_file = row.get("source_file")
    sheet = row.get("source_sheet")
    source_row = row.get("source_row")
    if not source_file or not sheet or not source_row:
        return []
    try:
        wb = load_workbook(ROOT / str(source_file), data_only=True)
        ws = wb[str(sheet)]
        values = [cell.value for cell in ws[int(source_row)]]
        wb.close()
        return values
    except Exception:
        return []


def wait_days(notes: str | None) -> int | None:
    if not notes:
        return None
    match = re.search(r"(\d+)\s*-\s*(\d+)\s*周", notes)
    if match:
        return int(match.group(2)) * 7
    match = re.search(r"(\d+)\s*周", notes)
    if match:
        return int(match.group(1)) * 7
    match = re.search(r"(\d+)\s*(?:天|日)", notes)
    if match:
        return int(match.group(1))
    match = re.search(r"(\d+)\s*个?\s*月", notes)
    if match:
        return int(match.group(1)) * 30
    return None


def calculate_confidence(row: dict[str, Any], candidate: dict[str, Any], source_evidence: str) -> float:
    score = 0.50
    notes = str(candidate.get("notes") or "")
    model = str(candidate.get("model") or "")
    trim = str(candidate.get("variant") or "")
    source_file = str(row.get("source_file") or "")
    raw_model = str(row.get("model") or "")
    raw_trim = str(row.get("variant") or row.get("trim_config") or "")
    evidence = f"{source_evidence} {notes}"

    if candidate.get("brand"):
        score += 0.10
    if candidate.get("model") and not re.search(r"\b(EXW|FCA|FOB|CIF|USD|CNY|RMB)\b|[$¥￥]\s*\d|\d{5,}", model, re.I):
        score += 0.15
    if any(candidate.get(field) not in (None, "", []) for field in ("supplier_price_cny", "cost_exw_usd", "cost_fca_usd", "cost_fob_usd")):
        score += 0.10
    if (
        ("CNY/RMB supplier price evidence" in evidence and candidate.get("supplier_price_cny"))
        or ("EXW USD evidence" in evidence and candidate.get("cost_exw_usd"))
        or ("FCA USD evidence" in evidence and candidate.get("cost_fca_usd"))
        or ("FOB USD evidence" in evidence and candidate.get("cost_fob_usd"))
    ):
        score += 0.10
    if candidate.get("supplier"):
        score += 0.10
    if candidate.get("location") or candidate.get("vehicle_supply_base"):
        score += 0.05
    if candidate.get("stock_quantity") and candidate.get("exterior_color"):
        score += 0.05
    if any(candidate.get(field) not in (None, "", []) for field in ("manufacture_year", "manufacture_month", "steering_setup", "market_region")):
        score += 0.05

    if candidate.get("supplier") and candidate.get("supplier") not in source_evidence and "# Path:" in source_evidence:
        score -= 0.10
    if model != raw_model or (trim and trim != raw_trim):
        score -= 0.10
    if any(candidate.get(field) not in (None, "", []) for field in ("supplier_price_cny", "cost_exw_usd", "cost_fca_usd", "cost_fob_usd")) and not re.search(r"\b(EXW|FCA|FOB|USD|CNY|RMB|人民币|美金|美元)\b|[$¥￥]", evidence, re.I):
        score -= 0.15
    if (candidate.get("location") or candidate.get("vehicle_supply_base")) and not row.get("location"):
        score -= 0.15
    if candidate.get("stock_quantity") and candidate.get("exterior_color") and re.search(r"[+＋]|\d+[^/]+/", str(row.get("exterior_color") or "")):
        score -= 0.15
    if any(token in source_evidence for token in ("merged:", "合并", "row None")) or row.get("source_row") in (None, ""):
        score -= 0.20
    if "霍尔果斯" in evidence and "南沙" in evidence and (candidate.get("location") or candidate.get("vehicle_supply_base")):
        score -= 0.30
    if re.search(r"\b[A-Z]{2,}\d{3,}|\d{5,}", raw_model + " " + raw_trim):
        score -= 0.30

    if any(token in evidence for token in ("地点冲突", "币种不明", "车型无法切分")):
        score = min(score, 0.59)
    return round(max(0.0, min(1.0, score)), 2)


def convert(row: dict[str, Any]) -> dict[str, Any] | None:
    brand = clean(row.get("brand"))
    values = source_values(row)
    raw_model = row.get("model")
    raw_trim = row.get("variant") or row.get("trim_config")
    source_file = str(row.get("source_file") or "")
    brand = brand if brand in VALID_BRANDS else path_brand_hint(source_file) or brand
    path_hint = path_model_hint(source_file, brand)
    supplier = clean(row.get("supplier"))
    supplier_note: str | None = None
    if supplier and supplier.startswith("主机厂") and source_file.startswith("APT\\"):
        supplier_note = f"source_supplier_text={supplier}"
        supplier = "APT"
    if brand == "Wuling" and values:
        raw_model = path_hint or raw_model
        raw_trim = " | ".join(clean(value) or "" for value in values[:5] if clean(value)) or raw_trim
    elif brand == "Changan" and values and clean(values[0]):
        raw_model = path_hint or values[0]
        raw_trim = " | ".join(clean(value) or "" for value in values[1:3] if clean(value)) or raw_trim
    elif brand == "Farizon" and ("新V系列" in source_file or path_hint) and len(values) >= 3:
        raw_model = path_hint or values[1] or raw_model
        raw_trim = clean(values[2]) or raw_trim
    elif brand == "Geely" and (path_hint or (values and clean(values[0]))):
        raw_model = path_hint or values[0]
        raw_trim = raw_trim or (clean(values[1]) if len(values) > 1 else clean(row.get("model")))
    elif brand in {"AVATR", "IM Motors", "Bestune", "Xiaomi"} and path_hint:
        raw_model = path_hint
    elif brand == "BYD" and values and any("海鸥" in str(value or "") for value in values[:3]):
        raw_model = next(value for value in values[:3] if "海鸥" in str(value or ""))
        raw_trim = " | ".join(clean(value) or "" for value in values[4:7] if clean(value)) or raw_trim
    elif brand == "BYD" and path_hint and re.search(r"(Sealion|Seagull|Yuan|Qin|Tang|Song|Seal|Shark|ATTO|BYD E7)", path_hint, re.I):
        raw_model = path_hint
    model, trim = split_model_trim(brand, raw_model, raw_trim)
    if any(re.search(r"(?:^|\s)选装[-_：:]", str(clean(value) or "")) for value in (raw_model, raw_trim, trim, source_file)):
        return None
    trim = normalize_variant_text(brand, model, trim)
    if clean(raw_model) in SUBTOTAL_TOKENS or clean(raw_trim) in SUBTOTAL_TOKENS or model in SUBTOTAL_TOKENS or trim in SUBTOTAL_TOKENS:
        return None
    if not brand or not model:
        return None
    if any(token in model for token in ("若选择", "报价", "备注", "说明")):
        return None

    qty, exterior = split_color_quantity(row.get("exterior_color"), row.get("stock_quantity"))
    year, month = parse_year_month(row.get("manufacture_date") or row.get("production_date"))
    location, vehicle_supply_base = split_place_fields(row.get("location"))
    source_evidence = evidence_text(row)
    business_notes = " | ".join(part for part in (clean(row.get("notes")), supplier_note) if part)
    internal_evidence_parts = []
    for key in ("cost_exw_cny", "cost_fca_cny", "cost_fob_cny", "official_suggested_price_cny"):
        if row.get(key) not in (None, ""):
            internal_evidence_parts.append(f"{key}={row.get(key)}")
    internal_evidence_parts.append(f"source={row.get('source_file')} sheet={row.get('source_sheet')} row={row.get('source_row')}")
    internal_evidence = " | ".join(part for part in internal_evidence_parts if part)
    supplier_cny = row.get("cost_fca_cny") or row.get("cost_fob_cny") or row.get("cost_exw_cny") or row.get("official_suggested_price_cny")
    candidate = {field: None for field in CURRENT_FIELDS}
    candidate.update(
        {
            "confidence": None,
            "brand": brand,
            "supplier": supplier,
            "model": model,
            "variant": trim,
            "supplier_price_cny": supplier_cny,
            "cost_fca_usd": row.get("cost_fca_usd"),
            "cost_fob_usd": row.get("cost_fob_usd"),
            "cost_exw_usd": row.get("cost_exw_usd"),
            "notes": business_notes,
            "order_wait_days": wait_days(business_notes),
            "manufacture_year": year,
            "manufacture_month": month,
            "min_quantity": row.get("min_quantity") or 1,
            "record_id": row.get("content_hash"),
            "interior_color": clean(row.get("interior_color")),
            "exterior_color": exterior,
            "location": location,
            "vehicle_supply_base": vehicle_supply_base,
            "steering_setup": clean(row.get("steering_setup")),
            "market_region": normalize_market(row.get("version_type")),
            "display_price_low": None,
            "stock_quantity": qty,
        }
    )
    if not any(candidate.get(field) not in (None, "", []) for field in ("cost_fca_usd", "cost_fob_usd", "cost_exw_usd", "supplier_price_cny", "stock_quantity")):
        return None
    confidence_evidence = f"{source_evidence}; {internal_evidence}"
    candidate["confidence"] = calculate_confidence(row, candidate, confidence_evidence)
    evidence = {field: source_evidence for field, value in candidate.items() if value not in (None, "", [])}
    if candidate.get("supplier_price_cny"):
        evidence["supplier_price_cny"] = f"{source_evidence}; {internal_evidence}; CNY/RMB supplier price evidence"
    for price_field, term in (("cost_exw_usd", "EXW"), ("cost_fca_usd", "FCA"), ("cost_fob_usd", "FOB")):
        if candidate.get(price_field):
            evidence[price_field] = f"{source_evidence}; {term} USD evidence"
    if candidate.get("stock_quantity") and candidate.get("exterior_color") and "/" not in str(candidate.get("exterior_color")):
        combo = f"{candidate['stock_quantity']}{candidate['exterior_color']}"
        if candidate.get("interior_color"):
            combo = f"{combo}/{candidate['interior_color']}"
        combo_evidence = f"# Path: {row.get('source_file')}; Sheet: {row.get('source_sheet')}; row {row.get('source_row')}; {combo}"
        evidence["stock_quantity"] = combo_evidence
        evidence["exterior_color"] = combo_evidence
        if candidate.get("interior_color"):
            evidence["interior_color"] = combo_evidence
    if candidate.get("supplier"):
        evidence["supplier"] = f"# Path: {row.get('source_file')}"
    if candidate.get("brand"):
        evidence["brand"] = f"# Path: {row.get('source_file')}; {source_evidence}"
    if candidate.get("model"):
        evidence["model"] = source_evidence
    candidate["_evidence"] = evidence
    candidate["_source_file"] = row.get("source_file")
    return candidate


def main() -> int:
    global ROOT
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", default="output/final/manual_feishu_ready_20260724.json")
    parser.add_argument("--output", default="output/final/current_feishu_candidates_20260727.json")
    parser.add_argument("--source-summary", default=None)
    parser.add_argument("--quality-prefix", default=None)
    parser.add_argument("--root", default=str(ROOT))
    args = parser.parse_args()
    ROOT = Path(args.root)
    rows = json.loads(Path(args.input).read_text("utf-8"))
    converted = [candidate for row in rows if (candidate := convert(row))]
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps({"records": converted}, ensure_ascii=False, indent=2), "utf-8")
    print(f"input_rows={len(rows)} output_records={len(converted)}")
    print(out)
    summary = {}
    if args.source_summary and Path(args.source_summary).exists():
        summary = json.loads(Path(args.source_summary).read_text("utf-8"))
    quality_prefix = Path(args.quality_prefix) if args.quality_prefix else out.with_suffix("")
    report = build_quality_report(converted, summary=summary, raw_records=rows)
    paths = write_report_files(report, quality_prefix)
    print("quality_report=" + paths["quality_report_json"])
    print("zero_row_sources=" + paths["zero_row_sources"])
    print("duplicate_candidates=" + paths["duplicate_candidates"])
    print("variant_audit=" + paths["variant_audit"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
