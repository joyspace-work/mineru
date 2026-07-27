from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path("C:/Users/HP/Downloads/\u8f66\u6e90\u6c47\u603b\u7eaf\u51c0 (2)/\u8f66\u6e90\u6c47\u603b\u7eaf\u51c0")
EXCLUDE = "\u7eaf\u51c0\u6c47\u603b.xlsx"
OUT = Path("output/final/manual_feishu_ready_20260724.json")
SUMMARY = Path("output/final/manual_feishu_ready_20260724_summary.json")

BRAND_ALIASES = {
    "BYD": "BYD",
    "比亚迪": "BYD",
    "吉利": "Geely",
    "远程": "Farizon",
    "东风": "Dongfeng",
    "五菱": "Wuling",
    "长安": "Changan",
    "深蓝": "Deepal",
    "DEEPAL": "Deepal",
    "小米": "Xiaomi",
    "广汽": "GAC Aion",
    "广汽埃安": "GAC Aion",
    "智己": "IM Motors",
    "阿维塔": "AVATR",
    "捷途": "Jetour",
    "福田": "Foton",
    "丰田": "Toyota",
    "零跑": "Leapmotor",
    "奔腾小马": "Bestune",
    "奔腾": "Bestune",
    "Smart": "Smart",
    "smart": "Smart",
    "福瑞通": "Farizon",
    "河南新商筹": "Farizon",
    "福瑞通4S店": "Farizon",
}


def clean(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = str(value).strip()
    return text if text not in {"", "/", "-", "—"} else None


def contains_any(text: str, tokens: list[str]) -> bool:
    lower = text.lower()
    return any(token.lower() in lower for token in tokens)


def number(value: Any) -> float | int | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else float(value)
    match = re.search(r"-?\d+(?:\.\d+)?", str(value).replace(",", ""))
    if not match:
        return None
    parsed = float(match.group(0))
    return int(parsed) if parsed.is_integer() else parsed


def normalize_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    match = re.search(r"(20\d{2})[年./-]\s*(\d{1,2})(?:[月./-]\s*(\d{1,2}))?", str(value))
    if not match:
        return None
    year, month, day = match.group(1), int(match.group(2)), int(match.group(3) or 1)
    return f"{year}-{month:02d}-{day:02d}"


def infer_supplier_brand(path: Path) -> tuple[str | None, str | None]:
    parts = path.relative_to(ROOT).parts
    supplier = parts[0] if parts else None
    brand = BRAND_ALIASES.get(supplier or "")
    for part in parts[1:]:
        if part in BRAND_ALIASES:
            brand = BRAND_ALIASES[part]
    for alias, mapped in BRAND_ALIASES.items():
        if alias.lower() in path.stem.lower():
            brand = mapped
    return supplier, brand


def normalize_brand(brand: Any, model: Any = None) -> str | None:
    text = clean(brand)
    if text in BRAND_ALIASES:
        return BRAND_ALIASES[text]
    if text:
        return text
    model_text = clean(model) or ""
    for alias, mapped in BRAND_ALIASES.items():
        if alias.lower() in model_text.lower():
            return mapped
    return None


def source_row(row: dict[str, Any], path: Path, sheet: str, row_number: int | None = None) -> dict[str, Any]:
    row["source_file"] = str(path.relative_to(ROOT))
    row["source_sheet"] = sheet
    if row_number is not None:
        row["source_row"] = row_number
    identity = "|".join(str(row.get(key) or "") for key in ("source_file", "source_sheet", "source_row", "model", "trim_config"))
    row["content_hash"] = hashlib.sha256(identity.encode("utf-8")).hexdigest()
    return row


def notes(*values: Any) -> str | None:
    result: list[str] = []
    for value in values:
        text = clean(value)
        if text and text not in result:
            result.append(text)
    return " | ".join(result) if result else None


def sheet_rows(ws) -> list[list[Any]]:
    return [[cell.value for cell in row] for row in ws.iter_rows()]


def header_score(row: list[Any]) -> int:
    text = " ".join(clean(value) or "" for value in row)
    tokens = ["品牌", "车型", "车系", "型号", "model", "config", "配置", "版本", "指导价", "msrp", "domestic price", "self-pickup", "售价", "报价", "价格", "exw", "fca", "fob", "数量", "颜色", "内饰", "外观"]
    return sum(1 for token in tokens if token.lower() in text.lower())


def find_header(rows: list[list[Any]]) -> int | None:
    best_index = None
    best_score = 0
    for index, row in enumerate(rows[:12]):
        score = header_score(row)
        if score > best_score:
            best_index = index
            best_score = score
    return best_index if best_score >= 2 else None


def column_kind(header: str, context: str) -> str | None:
    text = f"{context} {header}"
    header_lower = header.lower()
    if contains_any(header, ["品牌", "brand"]):
        return "brand"
    if contains_any(header, ["供应商", "车源供应方", "经销商"]):
        return "supplier"
    if "exw" in header_lower:
        return "cost_exw_usd" if contains_any(header, ["usd", "美金", "美元", "$"]) else "cost_exw_cny"
    if "fca" in header_lower:
        return "cost_fca_usd" if contains_any(header, ["usd", "美金", "美元", "$"]) else "cost_fca_cny"
    if "fob" in header_lower:
        return "cost_fob_usd" if contains_any(header, ["usd", "美金", "美元", "$"]) else "cost_fob_cny"
    if contains_any(header, ["车型名称", "车辆名称", "车型及配置信息", "car model", "型号 / model", "model name", "车系", "系列 / series", "车型", "型号"]):
        return "model"
    if contains_any(header, ["配置", "版本", "款型", "config", "vehicle series", "车型版本", "动力+座椅布局", "车辆类型", "电池", "发动机代码", "车型代码", "动力engine"]):
        return "trim_config"
    if contains_any(header, ["生产日期", "生产年月", "出厂", "production"]):
        return "manufacture_date"
    if contains_any(header, ["数量", "库存", "可售数量", "合计", "stock", "qty"]):
        return "stock_quantity"
    if contains_any(header, ["外观", "外饰", "exterior", "颜色"]):
        return "exterior_color"
    if contains_any(header, ["内饰", "interior"]):
        return "interior_color"
    if contains_any(header, ["车源地", "所在地", "提货", "交付地点", "地点", "location", "exw地点"]):
        return "location"
    if contains_any(header, ["左右", "方向", "舵"]):
        return "steering_setup"
    if contains_any(header, ["备注", "说明", "notes", "remark", "选配", "主要配置描述", "major configurations", "出口方式", "销售方式", "状态"]):
        return "notes"
    if contains_any(header, ["指导价", "msrp", "domestic", "国内价格", "售价", "市场指导价"]):
        return "official_suggested_price_cny"
    if contains_any(header, ["报价", "价格", "人民币 rmb", "美金usd", "美元 usd"]):
        return "cost_exw_usd" if contains_any(text, ["usd", "美金", "美元", "$"]) else "cost_exw_cny"
    return None


def split_color_quantity(value: Any) -> list[tuple[float | int | None, str | None, str | None]]:
    text = clean(value)
    if not text:
        return []
    result = []
    for part in re.split(r"[+＋]", text):
        part = part.strip()
        match = re.match(r"(?:(\d+)\s*)?([^/（）()\d]+)(?:[/（(]([^）)]+)[）)]?)?$", part)
        if match:
            result.append((number(match.group(1)), clean(match.group(2)), clean(match.group(3))))
        else:
            result.append((None, part, None))
    return result


def parse_text_like(path: Path, sheet: str, rows: list[list[Any]]) -> list[dict[str, Any]]:
    supplier, _brand_hint = infer_supplier_brand(path)
    texts = [clean(row[1] if len(row) > 1 else None) for row in rows[1:]]
    texts = [text for text in texts if text]
    joined = "\n".join(texts)
    rel = str(path.relative_to(ROOT))
    output = []
    if "Smart #" in joined:
        location = "宁波/上海" if ("宁波" in joined or "上海" in joined) else None
        for text in texts:
            match = re.search(r"(Smart\s*#\d)\s+(\w+)\s*\$?([0-9,]+)", text, re.I)
            if match:
                output.append(source_row({"brand": "Smart", "model": match.group(1), "trim_config": match.group(2), "cost_fca_usd": number(match.group(3)), "location": location, "supplier": supplier}, path, sheet))
    elif "ATTO3" in rel or "超越型" in joined:
        for text in texts:
            match = re.search(r"([^\s的]+型).*?EXW\s*([^\d\s是:：]+)?\s*(?:是|:|：)?\s*\$?([0-9,]+)", text, re.I)
            if match:
                output.append(source_row({"brand": "BYD", "model": "ATTO3", "trim_config": match.group(1), "location": match.group(2), "cost_exw_usd": number(match.group(3)), "supplier": supplier}, path, sheet))
    elif "Global automotive" in rel:
        for text in texts:
            match = re.search(r"(.+?)\s+([0-9]{5,6})\s+EXW[:：]?\s*\$?([0-9,]+)", text, re.I)
            if match:
                output.append(source_row({"brand": "BYD", "model": match.group(1).strip(), "official_suggested_price_cny": number(match.group(2)), "cost_exw_usd": number(match.group(3)), "supplier": supplier}, path, sheet))
    elif "智己L6" in rel or "智己LS6" in rel:
        pending = None
        for text in texts:
            match = re.search(r"(L6|LS6)\s+(.+?)\s*指导价\s*([0-9,]+)", text, re.I)
            if match:
                pending = {"brand": "IM Motors", "model": match.group(1), "trim_config": match.group(2).strip(), "official_suggested_price_cny": number(match.group(3)), "supplier": supplier}
                continue
            price_match = re.search(r"([0-9,]+)\s+EXW\s*([^\s]+)", text, re.I)
            if pending and price_match:
                pending["cost_exw_usd"] = number(price_match.group(1))
                pending["location"] = price_match.group(2)
                output.append(source_row(pending, path, sheet))
                pending = None
    elif "i60" in rel:
        for text in texts:
            match = re.search(r"(现车)?\s*(\d+\S*)\s+EXW\s*([^\d\s]+)\s*([0-9,]+)\s*USD", text, re.I)
            if match:
                output.append(source_row({"brand": "GAC Aion", "model": "i60", "trim_config": match.group(2), "status_vehicle": "现车" if match.group(1) else None, "location": match.group(3), "cost_exw_usd": number(match.group(4)), "supplier": supplier}, path, sheet))
    return output


def parse_parameter_sheet(path: Path, sheet: str, rows: list[list[Any]]) -> list[dict[str, Any]]:
    supplier, brand_hint = infer_supplier_brand(path)
    kv = {clean(row[0]): clean(row[1]) for row in rows if len(row) >= 2 and clean(row[0])}
    if "EXW价格" in kv:
        model = path.stem.replace(f" EXW{kv.get('EXW价格')}", "")
        spec_notes = notes(*(f"{key}: {value}" for key, value in kv.items() if key != "EXW价格"))
        return [source_row({"brand": brand_hint or supplier, "model": model, "cost_exw_cny": number(kv["EXW价格"]), "supplier": supplier, "notes": spec_notes}, path, sheet, 2)]
    if "车辆名称" in kv:
        spec_notes = notes(*(f"{key}: {value}" for key, value in list(kv.items())[:20] if key != "车辆名称"))
        return [source_row({"brand": brand_hint or supplier, "model": kv["车辆名称"], "trim_config": sheet, "supplier": supplier, "notes": spec_notes}, path, sheet)]
    return []


def parse_table_sheet(path: Path, sheet: str, rows: list[list[Any]]) -> list[dict[str, Any]]:
    supplier_hint, brand_hint = infer_supplier_brand(path)
    header_index = find_header(rows)
    if header_index is None:
        return []
    context_rows = rows[max(0, header_index - 3):min(len(rows), header_index + 3)]
    context = " ".join(clean(value) or "" for row in context_rows for value in row)
    max_cols = max((len(row) for row in rows), default=0)
    headers = []
    for col in range(max_cols):
        parts = []
        for row_index in range(max(0, header_index - 2), min(len(rows), header_index + 3)):
            value = clean(rows[row_index][col]) if col < len(rows[row_index]) else None
            if value and value not in parts:
                parts.append(value)
        headers.append(" ".join(parts))
    kinds = [column_kind(header, context) for header in headers]
    if not any(kind in kinds for kind in ("model", "trim_config", "official_suggested_price_cny", "cost_exw_usd", "cost_fca_usd", "cost_fob_usd", "cost_exw_cny", "cost_fca_cny", "cost_fob_cny")):
        return []
    output = []
    last: dict[str, Any] = {}
    for row_number, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        if not any(clean(value) for value in row):
            continue
        record: dict[str, Any] = {}
        for col, kind in enumerate(kinds):
            if not kind or col >= len(row):
                continue
            value = clean(row[col])
            if kind in {"model", "brand", "supplier"}:
                if value:
                    last[kind] = value
                elif kind in last:
                    value = last[kind]
            if not value:
                continue
            if kind in {"official_suggested_price_cny", "cost_exw_usd", "cost_fca_usd", "cost_fob_usd", "cost_exw_cny", "cost_fca_cny", "cost_fob_cny", "stock_quantity"}:
                parsed = number(value)
                if parsed is not None:
                    record[kind] = parsed
            elif kind == "manufacture_date":
                record[kind] = normalize_date(value)
            elif kind == "location":
                record[kind] = re.sub(r"^(EXW|FCA|FOB)\s*", "", str(value), flags=re.I)
            elif kind == "steering_setup":
                raw = str(value).upper()
                record[kind] = "左舵" if ("左" in raw or "LHD" in raw) else ("右舵" if ("右" in raw or "RHD" in raw) else value)
            elif kind == "notes":
                record[kind] = notes(record.get(kind), value)
            elif kind == "trim_config" and record.get("trim_config"):
                record["trim_config"] = f"{record['trim_config']} {value}"
            else:
                record[kind] = value
        record.setdefault("supplier", supplier_hint)
        record["brand"] = normalize_brand(record.get("brand") or brand_hint, record.get("model"))
        if not record.get("brand") and (contains_any(str(record.get("model") or ""), ["V6E", "V7E", "V8E", "SV"]) or contains_any(str(record.get("trim_config") or ""), ["V6E", "V7E", "V8E", "SV"])):
            record["brand"] = "Farizon"
        if not record.get("brand") and contains_any(str(record.get("model") or ""), ["S05", "S07"]):
            record["brand"] = "Deepal"
        if not record.get("model") and record.get("trim_config"):
            record["model"] = record["trim_config"]
            record["trim_config"] = None
        if not record.get("model") or record.get("model") in {"小计", "合计"} or "小计" in str(record.get("model")):
            continue
        base = source_row(record, path, sheet, row_number)
        combos = split_color_quantity(base.get("exterior_color")) if base.get("exterior_color") and re.search(r"\d.*[+＋/]|[+＋]", str(base.get("exterior_color"))) else []
        if len(combos) > 1:
            for quantity, exterior, interior in combos:
                child = dict(base)
                if quantity is not None:
                    child["stock_quantity"] = quantity
                child["exterior_color"] = exterior
                if interior:
                    child["interior_color"] = interior
                output.append(child)
        else:
            output.append(base)
    return output


def meaningful(row: dict[str, Any]) -> bool:
    if not row.get("brand") or not row.get("model"):
        return False
    for key in ("cost_exw_usd", "cost_exw_cny", "cost_fca_usd", "cost_fca_cny", "cost_fob_usd", "cost_fob_cny", "official_suggested_price_cny", "stock_quantity", "notes"):
        if row.get(key) not in (None, ""):
            return True
    return False


def correct_currency_scale(row: dict[str, Any]) -> None:
    official = number(row.get("official_suggested_price_cny"))
    for term in ("exw", "fca", "fob"):
        usd_key = f"cost_{term}_usd"
        cny_key = f"cost_{term}_cny"
        value = number(row.get(usd_key))
        if value is None:
            continue
        if value > 50000 or (official and value > official * 0.5):
            row[cny_key] = row.get(cny_key) or value
            row[usd_key] = None


def repair_path_brand(row: dict[str, Any]) -> None:
    source = row.get("source_file")
    if not source:
        return
    _supplier, path_brand = infer_supplier_brand(ROOT / str(source))
    if not path_brand:
        return
    known_brands = set(BRAND_ALIASES.values())
    current = row.get("brand")
    if current and current not in known_brands and current != path_brand:
        old_model = row.get("model")
        row["brand"] = path_brand
        row["model"] = current
        if old_model and old_model != current:
            row["trim_config"] = notes(old_model, row.get("trim_config"))


def main() -> int:
    rows: list[dict[str, Any]] = []
    per_file = []
    files = sorted(path for path in ROOT.rglob("*.xlsx") if path.name != EXCLUDE)
    for path in files:
        workbook = load_workbook(path, data_only=True)
        file_rows: list[dict[str, Any]] = []
        for ws in workbook.worksheets:
            if "来源" in ws.title or "费用" in ws.title or ws.max_row <= 1:
                continue
            values = sheet_rows(ws)
            header = " ".join(clean(value) or "" for value in (values[0] if values else []))
            parsed: list[dict[str, Any]] = []
            if "内容" in header and "拆分" in header:
                parsed = parse_text_like(path, ws.title, values)
            if not parsed:
                parsed = parse_table_sheet(path, ws.title, values)
            if parsed and not any(meaningful(row) for row in parsed):
                parsed = []
            if not parsed:
                parsed = parse_parameter_sheet(path, ws.title, values)
            file_rows.extend(parsed)
        cleaned = []
        for row in file_rows:
            if not row.get("brand"):
                row["brand"] = normalize_brand(None, row.get("model"))
            raw = " ".join(str(row.get(key) or "") for key in ("notes", "trim_config", "model"))
            if not row.get("version_type"):
                if contains_any(raw, ["国际", "出口", "欧标", "海外"]):
                    row["version_type"] = "国际版"
                elif contains_any(raw, ["国内", "中规"]):
                    row["version_type"] = "国内版"
            if not row.get("steering_setup"):
                if "右舵" in raw or "RHD" in raw.upper():
                    row["steering_setup"] = "右舵"
                elif "左舵" in raw or "LHD" in raw.upper():
                    row["steering_setup"] = "左舵"
            repair_path_brand(row)
            correct_currency_scale(row)
            if meaningful(row):
                cleaned.append(row)
        rows.extend(cleaned)
        per_file.append({"file": str(path.relative_to(ROOT)), "rows": len(cleaned)})
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(rows, ensure_ascii=False, indent=2), "utf-8")
    SUMMARY.write_text(json.dumps({"source_root": str(ROOT), "excluded": EXCLUDE, "source_file_count": len(files), "candidate_count": len(rows), "per_file": per_file}, ensure_ascii=False, indent=2), "utf-8")
    print(f"source_file_count={len(files)}")
    print(f"candidate_count={len(rows)}")
    print(OUT)
    print(SUMMARY)
    print(json.dumps(per_file, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
