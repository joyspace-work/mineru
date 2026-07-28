from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


TARGET_FIELDS = [
    "model_id", "brand", "model", "trim_config", "production_date",
    "exterior_color", "interior_color", "stock_quantity", "supplier", "location", "notes",
    "min_quantity", "max_quantity", "official_suggested_price_cny",
    "cost_fca_usd", "cost_fob_usd", "cost_exw_usd", "steering_setup",
    "order_wait_days", "market_region",
]


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def clean_number(value: Any) -> float | int | None:
    raw = text(value)
    if not raw:
        return None
    match = re.search(r"[-+]?\d[\d,]*(?:\.\d+)?", raw)
    if not match:
        return None
    number = float(match.group(0).replace(",", ""))
    return int(number) if number.is_integer() else number


def currency_number(value: Any) -> float | int | None:
    raw = text(value)
    match = re.search(r"(?:\$|USD\s*)\s*(\d[\d,]*(?:\.\d+)?)|(\d[\d,]*(?:\.\d+)?)\s*(?:USD|美元|美金)", raw, re.I)
    if not match:
        return clean_number(value)
    raw_number = (match.group(1) or match.group(2) or "").replace(",", "")
    number = float(raw_number)
    return int(number) if number.is_integer() else number


def date_value(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d 00:00:00")
    raw = text(value)
    if not raw:
        return None
    if re.fullmatch(r"20\d{2}", raw):
        return f"{raw}-01-01 00:00:00"
    m = re.search(r"(20\d{2})\D+(\d{1,2})(?:\D+(\d{1,2}))?", raw)
    if m:
        day = m.group(3) or "1"
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(day):02d} 00:00:00"
    return None


def normalize_location(value: Any) -> str | None:
    raw = text(value)
    if not raw:
        return None
    raw = re.sub(r"^(EXW|FCA|FOB|CIF)\s*", "", raw, flags=re.I)
    raw = raw.replace("Horgos", "霍尔果斯").replace("Khorgos", "霍尔果斯")
    return raw.strip(" ：:")


def path_parts(path: Path, root: Path) -> list[str]:
    return list(path.relative_to(root).parts)


def supplier_from_path(path: Path, root: Path) -> str:
    parts = path_parts(path, root)
    return parts[0] if parts else ""


def brand_from_path(path: Path, root: Path) -> str | None:
    parts = path_parts(path, root)
    known = {
        "BYD": "BYD", "比亚迪": "BYD", "吉利": "Geely", "五菱": "Wuling",
        "长安": "Changan", "东风": "Dongfeng", "广汽": "GAC Aion",
        "丰田": "Toyota", "智己": "IM", "捷途": "Jetour", "Smart": "Smart",
        "上汽": "SAIC", "小米": "Xiaomi", "福田": "Foton",
    }
    for part in parts:
        if part in known:
            return known[part]
    return None


def infer_model_from_path(path: Path) -> str | None:
    stem = path.stem
    for token in ("价格表", "报价", "价格", "库存", "EXW", "FOB"):
        stem = stem.replace(token, "")
    stem = re.sub(r"\d+(?:\.\d+)?$", "", stem).strip(" _-、")
    return stem or None


def split_color(value: Any) -> tuple[str | None, str | None]:
    raw = text(value)
    if not raw:
        return None, None
    m = re.match(r"(.+?)\((.+)\)$", raw)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    if "/" in raw and len(raw.split("/", 1)) == 2:
        a, b = raw.split("/", 1)
        return a.strip(), b.strip()
    return raw, None


def add_note(existing: str | None, extra: str | None) -> str | None:
    extra = text(extra)
    if not extra:
        return existing
    if existing:
        if extra in existing:
            return existing
        return f"{existing} | {extra}"
    return extra


def is_headerish(row: list[Any]) -> bool:
    exact = {
        "品牌", "车系", "系列", "车型", "型号", "配置", "版本", "指导价", "价格", "报价",
        "数量", "库存", "合计", "售价", "备注", "外观", "内饰", "颜色", "车源地",
        "生产日期", "生产年月", "市场指导价",
    }
    hits = 0
    for value in row:
        label = text(value)
        low = label.lower()
        compact = re.sub(r"\s+", "", label)
        if compact in exact:
            hits += 1
        elif compact.lower() in {"车型model", "车型/型号", "carmodel车型", "fob", "fca", "exw", "config"}:
            hits += 1
        elif any(token in low for token in ["msrp", "modelcode", "model code", "domestic price", "guide price"]):
            hits += 1
        elif any(token in low for token in ["fob", "fca", "exw"]) and any(token in low for token in ["price", "rmb", "usd", "人民币", "美金", "美元", "报价"]):
            hits += 1
        elif "/" in label and any(token in low for token in ["series", "model", "price", "usd", "rmb"]):
            hits += 1
    return hits >= 2


def header_map(row: list[Any]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    labels = [text(v) for v in row]
    has_model_col = any(("车型" in s or "Car Model" in s or "系列" in s) for s in labels)
    for i, label in enumerate(labels):
        low = label.lower()
        if not label:
            continue
        if "品牌" in label or label.strip().lower() == "brand":
            mapping[i] = "brand"
        elif "供应" in label or "车源供应" in label:
            mapping[i] = "supplier"
        elif "车型代码" in label or "modelcode" in low or "发动机代码" in label or "engine code" in low:
            mapping[i] = "notes"
        elif low.strip() == "config" and i == 0:
            mapping[i] = "model"
        elif "配置" in label or "版本" in label or "车辆类型" in label or low.strip() in {"config", "configuration"}:
            mapping[i] = "trim_config"
        elif "型号" in label or "model" in low:
            mapping[i] = "trim_config" if has_model_col else "model"
        elif "车系" in label or "车型" in label or "series" in low or "car model" in low:
            mapping[i] = "model"
        elif "指导价" in label or "msrp" in low or "domestic price" in low or "市场指导价" in label:
            mapping[i] = "official_suggested_price_cny"
        elif label.strip() in {"售价", "国内价格"}:
            mapping[i] = "official_suggested_price_cny"
        elif "外观" in label or "exterior" in low:
            mapping[i] = "exterior_color"
        elif "内饰" in label or "interior" in low:
            mapping[i] = "interior_color"
        elif "颜色" in label or low == "color":
            mapping[i] = "color"
        elif "数量" in label or "库存" in label or "台数" in label or label.strip() == "合计":
            mapping[i] = "stock_quantity"
        elif "生产" in label or "年月" in label:
            mapping[i] = "production_date"
        elif "车源地" in label or "交付地点" in label or "地点" in label or "仓库" in label:
            mapping[i] = "location"
        elif "备注" in label or "说明" in label or "major configurations" in low:
            mapping[i] = "notes"
        elif "报价" in label and "人民币" in label:
            mapping[i] = "cost_exw_rmb_note"
        elif "报价" in label and ("usd" in low or "美金" in label or "美元" in label):
            mapping[i] = "cost_exw_usd"
        elif label.strip() in {"报价", "价格"}:
            mapping[i] = "cost_exw_auto"
        elif "self-pickup" in low and "usd" in low:
            mapping[i] = "cost_fca_usd"
        elif ("khorgos" in low or "horgos" in low) and "usd" in low:
            mapping[i] = "cost_fca_usd"
        elif "exw" in low and ("usd" in low or "美金" in label or "美元" in label):
            mapping[i] = "cost_exw_usd"
        elif "fca" in low and ("usd" in low or "美金" in label or "美元" in label):
            mapping[i] = "cost_fca_usd"
        elif "fob" in low and ("usd" in low or "美金" in label or "美元" in label):
            mapping[i] = "cost_fob_usd"
        elif "exw" in low and ("人民币" in label or "rmb" in low or "cny" in low):
            mapping[i] = "cost_exw_rmb_note"
        elif "fca" in low and ("人民币" in label or "rmb" in low or "cny" in low):
            mapping[i] = "cost_fca_rmb_note"
        elif "fob" in low and ("人民币" in label or "rmb" in low or "cny" in low):
            mapping[i] = "cost_fob_rmb_note"
        elif "exw" in low:
            mapping[i] = "cost_exw_auto"
        elif "fca" in low:
            mapping[i] = "cost_fca_auto"
        elif "fob" in low:
            mapping[i] = "cost_fob_auto"
    return mapping


def apply_auto_price(record: dict[str, Any], key: str, value: Any, label: str) -> None:
    number = clean_number(value)
    if number is None:
        return
    if number < 50000:
        target = key.replace("_auto", "_usd")
        record[target] = number
    else:
        record["notes"] = add_note(record.get("notes"), f"{label}:{number}")


def normalize_brand_model(record: dict[str, Any]) -> None:
    brand = text(record.get("brand"))
    model = text(record.get("model"))
    if brand in {"比亚迪"}:
        record["brand"] = "BYD"
    elif brand in {"吉利"}:
        record["brand"] = "Geely"
    elif brand in {"五菱"}:
        record["brand"] = "Wuling"
    elif brand in {"长安"}:
        record["brand"] = "Changan"
    elif brand in {"东风"}:
        record["brand"] = "Dongfeng"
    elif brand in {"广汽埃安", "广汽"}:
        record["brand"] = "GAC Aion"
    elif brand in {"零跑"}:
        record["brand"] = "Leapmotor"
    elif brand in {"小米"}:
        record["brand"] = "Xiaomi"
    elif brand in {"智己"}:
        record["brand"] = "IM"
    elif brand in {"丰田"}:
        record["brand"] = "Toyota"
    elif brand in {"福田"}:
        record["brand"] = "Foton"
    if not record.get("brand"):
        low = model.lower()
        if "byd" in low or any(x in model for x in ["海狮", "元PLUS", "元UP", "秦", "海鸥", "唐L", "鲨鱼", "铁3"]):
            record["brand"] = "BYD"
        elif "geely" in low or "吉利" in model or "银河" in model:
            record["brand"] = "Geely"
        elif "wuling" in low or "五菱" in model or "星光" in model:
            record["brand"] = "Wuling"
        elif "deepal" in low or model in {"S05", "S07", "L06", "S09"}:
            record["brand"] = "Deepal"
        elif "smart" in low:
            record["brand"] = "Smart"
        elif "t03" in low:
            record["brand"] = "Leapmotor"
        elif any(token in model for token in ["V6E", "V7E", "V8E", "SV", "星享", "行享", "行镖", "中轴", "长轴", "短轴"]):
            record["brand"] = "Farizon"
    if record.get("brand") == "BYD":
        replacements = {"元UP": "Yuan UP", "海鸥": "Seagull", "海狮07": "Sealion 7", "海狮 07": "Sealion 7"}
        for raw, canonical in replacements.items():
            if raw in model:
                record["model"] = model.replace(raw, canonical)
                break


def finalize(record: dict[str, Any], source: Path, root: Path) -> dict[str, Any] | None:
    if record.get("color"):
        exterior, interior = split_color(record.pop("color"))
        record.setdefault("exterior_color", exterior)
        record.setdefault("interior_color", interior)
    for key in ["cost_exw_auto", "cost_fca_auto", "cost_fob_auto"]:
        if key in record:
            apply_auto_price(record, key, record.pop(key), key)
    for key in list(record):
        if key.endswith("_rmb_note"):
            record["notes"] = add_note(record.get("notes"), f"{key.replace('_note','')}:{record.pop(key)}")
    if record.get("production_date"):
        record["production_date"] = date_value(record["production_date"])
    if record.get("location"):
        record["location"] = normalize_location(record["location"])
    if record.get("steering_setup") in {"左", "LHD"}:
        record["steering_setup"] = "左舵"
    elif record.get("steering_setup") in {"右", "RHD"}:
        record["steering_setup"] = "右舵"
    path_brand = brand_from_path(source, root)
    normalize_brand_model(record)
    canonical = {"BYD", "Geely", "Wuling", "Changan", "Dongfeng", "GAC Aion", "Toyota", "IM", "Jetour", "Smart", "SAIC", "Xiaomi", "Foton", "Leapmotor", "Deepal"}
    if path_brand and record.get("brand") not in canonical:
        record["brand"] = path_brand
    record.setdefault("supplier", supplier_from_path(source, root))
    record.setdefault("brand", path_brand)
    model = text(record.get("model"))
    trim = text(record.get("trim_config"))
    if not model and trim:
        record["model"] = trim
        record["trim_config"] = None
    if not text(record.get("model")):
        return None
    if not any(record.get(k) not in (None, "") for k in ["official_suggested_price_cny", "cost_fca_usd", "cost_fob_usd", "cost_exw_usd", "stock_quantity"]):
        # Keep sparse inventory rows when they still carry a trim and supplier, but skip pure text notes.
        if not trim:
            return None
    record["notes"] = add_note(record.get("notes"), f"source={source.relative_to(root)}")
    return {field: record.get(field) for field in TARGET_FIELDS}


def parse_table_sheet(path: Path, root: Path, ws: Any) -> list[dict[str, Any]]:
    rows = [[cell.value for cell in row] for row in ws.iter_rows()]
    out: list[dict[str, Any]] = []
    current: dict[str, Any] = {}
    header: dict[int, str] = {}
    header_labels: dict[int, str] = {}
    for row in rows:
        if is_headerish(row):
            mapped = header_map(row)
            has_identity = any(field in {"brand", "model", "trim_config"} for field in mapped.values())
            if header and not has_identity:
                header.update(mapped)
                for i, value in enumerate(row):
                    if text(value):
                        header_labels[i] = add_note(header_labels.get(i), text(value)) or text(value)
                continue
            if header and has_identity:
                for col, field in header.items():
                    if col not in mapped and ("cost_" in field or field in {"official_suggested_price_cny", "stock_quantity"}):
                        mapped[col] = field
            header = mapped
            header_labels = {i: text(v) for i, v in enumerate(row)}
            current = {}
            continue
        if not header:
            continue
        if not any(text(v) for v in row):
            continue
        rec = dict(current)
        for i, key in header.items():
            if i >= len(row):
                continue
            val = row[i]
            if text(val) == "":
                continue
            if key in {"brand", "model", "trim_config", "location", "supplier"}:
                rec[key] = text(val)
                current[key] = rec[key]
            elif key == "notes":
                rec["notes"] = add_note(rec.get("notes"), text(val))
            elif key == "color":
                rec["color"] = text(val)
            elif key in {"official_suggested_price_cny", "stock_quantity", "cost_exw_usd", "cost_fca_usd", "cost_fob_usd"}:
                rec[key] = clean_number(val)
            elif key in {"cost_exw_auto", "cost_fca_auto", "cost_fob_auto"}:
                rec[key] = val
            elif key.endswith("_rmb_note"):
                rec[key] = val
            elif key == "production_date":
                rec[key] = val
            elif key == "steering_setup":
                rec[key] = text(val)
        if "左右" in " ".join(header_labels.values()) and len(row) > 5 and text(row[5]) in {"左", "右", "LHD", "RHD"}:
            rec["steering_setup"] = text(row[5])
        if header.get(0) == "model" and 1 not in header and len(row) > 1 and text(row[1]):
            rec.setdefault("trim_config", text(row[1]))
        if any(k in rec for k in ["model", "trim_config"]):
            final = finalize(rec, path, root)
            if final:
                out.append(final)
    return out


def parse_key_value_sheet(path: Path, root: Path, ws: Any) -> list[dict[str, Any]]:
    pairs = [(text(r[0].value), r[1].value if len(r) > 1 else None, text(r[2].value) if len(r) > 2 else "") for r in ws.iter_rows()]
    has_price = any(k in {"EXW价格", "FOB价格", "FCA价格"} for k, _, _ in pairs)
    has_vehicle_name = any(k == "车辆名称" for k, _, _ in pairs)
    if not has_price and not has_vehicle_name:
        return []
    rec: dict[str, Any] = {
        "supplier": supplier_from_path(path, root),
        "brand": brand_from_path(path, root),
        "model": infer_model_from_path(path),
    }
    notes = []
    for key, value, remark in pairs:
        if not key:
            continue
        if key == "车辆名称":
            rec["model"] = text(value)
            rec["trim_config"] = ws.title
        elif key == "EXW价格":
            rec["cost_exw_auto"] = value
        elif key == "FOB价格":
            rec["cost_fob_auto"] = value
        elif key == "FCA价格":
            rec["cost_fca_auto"] = value
        elif key == "颜色":
            rec["exterior_color"] = text(value)
        elif key not in {"参数类别"}:
            notes.append(f"{key}:{text(value)}" + (f" ({remark})" if remark else ""))
    rec["notes"] = " | ".join(notes[:12])
    final = finalize(rec, path, root)
    return [final] if final else []


def parse_text_sheet(path: Path, root: Path, ws: Any) -> list[dict[str, Any]]:
    lines = [text(row[1].value or row[2].value or row[0].value) for row in ws.iter_rows(min_row=2)]
    lines = [line for line in lines if line]
    if not lines:
        return []
    out: list[dict[str, Any]] = []
    supplier = supplier_from_path(path, root)
    path_brand = brand_from_path(path, root)
    default_model = infer_model_from_path(path)
    location_hint = None
    for line in lines:
        m_loc = re.search(r"(FCA|FOB|EXW)\s*([A-Za-z\u4e00-\u9fff/]+)", line, re.I)
        if "以上报价" in line and m_loc:
            location_hint = normalize_location(m_loc.group(0))
    for line in lines:
        m_loc = re.search(r"(FCA|FOB|EXW)\s*([A-Za-z\u4e00-\u9fff/]+)", line, re.I)
        if "以上报价" in line and m_loc:
            continue
        if "选配" in line or "冰箱" in line or "激光雷达" in line or "轮毂" in line:
            continue
        rec: dict[str, Any] = {"supplier": supplier, "brand": path_brand, "model": default_model, "notes": None}
        if "$" in line or re.search(r"\bUSD\b", line, re.I):
            price = currency_number(line)
            term = "cost_fca_usd" if ("FCA" in line.upper() or location_hint) else "cost_fob_usd" if "FOB" in line.upper() else "cost_exw_usd"
            rec[term] = price
            rec["model"] = re.sub(r"\$[\d,]+|\bUSD\b|EXW|FCA|FOB|南沙|宁波|上海|[:：]", "", line, flags=re.I).strip()
        elif "EXW" in line.upper() or "FCA" in line.upper() or "FOB" in line.upper():
            price = clean_number(line)
            term = "cost_exw_usd" if "EXW" in line.upper() else "cost_fca_usd" if "FCA" in line.upper() else "cost_fob_usd"
            rec[term] = price
            rec["location"] = normalize_location(m_loc.group(0)) if m_loc else None
            model_text = re.split(r"EXW|FCA|FOB", line, flags=re.I)[0]
            model_text = re.sub(r"指导价\s*\d+|是|的", "", model_text).strip()
            if model_text:
                rec["trim_config"] = model_text if default_model and default_model not in model_text else None
                rec["model"] = default_model if default_model else model_text
        else:
            m = re.search(r"(.+?)\s+(\d[\d,]{3,})", line)
            if m:
                rec["model"] = m.group(1).strip()
                rec["cost_fca_usd"] = clean_number(m.group(2))
        if not rec.get("location"):
            rec["location"] = location_hint
        final = finalize(rec, path, root)
        if final:
            out.append(final)
    return out


def parse_file(path: Path, root: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, data_only=True, read_only=False)
    out: list[dict[str, Any]] = []
    try:
        for ws in wb.worksheets:
            if ws.title == "来源" or not any(text(c.value) for row in ws.iter_rows(max_row=5) for c in row):
                continue
            first_rows = [" ".join(text(c.value) for c in row) for row in ws.iter_rows(min_row=1, max_row=4)]
            joined = " ".join(first_rows)
            if "行号" in joined and "内容" in joined:
                out.extend(parse_text_sheet(path, root, ws))
            out.extend(parse_key_value_sheet(path, root, ws))
            out.extend(parse_table_sheet(path, root, ws))
    finally:
        wb.close()
    # De-duplicate identical rows from sheets with both summary and spec sheets.
    seen = set()
    deduped = []
    for row in out:
        key = json.dumps(row, ensure_ascii=False, sort_keys=True)
        if key not in seen:
            seen.add(key)
            deduped.append(row)
    return deduped


def build_lark_batch(rows: list[dict[str, Any]]) -> dict[str, Any]:
    return {"fields": TARGET_FIELDS, "rows": [[row.get(field) for field in TARGET_FIELDS] for row in rows]}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", required=True)
    parser.add_argument("--out", default="output/final/xlsx_no_ai_candidates.json")
    parser.add_argument("--batch", default="output/final/xlsx_no_ai_feishu_batch.json")
    args = parser.parse_args()
    root = Path(args.root)
    all_rows: list[dict[str, Any]] = []
    counts: dict[str, int] = {}
    for path in sorted(root.rglob("*.xlsx")):
        rows = parse_file(path, root)
        counts[str(path.relative_to(root))] = len(rows)
        all_rows.extend(rows)
    out_path = Path(args.out)
    batch_path = Path(args.batch)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps({"counts": counts, "rows": all_rows}, ensure_ascii=False, indent=2), "utf-8")
    batch_path.write_text(json.dumps(build_lark_batch(all_rows), ensure_ascii=False, indent=2), "utf-8")
    print(f"files={len(counts)} rows={len(all_rows)}")
    for name, count in counts.items():
        print(f"{count:4d}  {name}")
    print(out_path)
    print(batch_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
