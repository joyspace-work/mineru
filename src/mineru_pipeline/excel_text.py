from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import csv
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".csv"}


@dataclass(frozen=True)
class ExcelText:
    source: Path
    text: str
    sheet_count: int
    row_count: int


def iter_excel_files(source: str | Path) -> list[Path]:
    root = Path(source)
    if root.is_file():
        return [root] if root.suffix.lower() in EXCEL_SUFFIXES and not root.name.startswith("~$") else []
    if not root.exists():
        return []
    return sorted(
        path
        for path in root.rglob("*")
        if path.is_file() and path.suffix.lower() in EXCEL_SUFFIXES and not path.name.startswith("~$")
    )


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value).replace("\r", " ").replace("\n", " ").strip()


def _trim_cells(values: list[str]) -> list[str]:
    while values and values[-1] == "":
        values.pop()
    return values


def _render_csv(path: Path) -> ExcelText:
    lines = [f"# Source: {path.name}", "## Sheet: csv"]
    row_count = 0
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        for row_number, row in enumerate(csv.reader(handle), start=1):
            values = _trim_cells([cell_text(value) for value in row])
            if not values:
                continue
            row_count += 1
            lines.append(f"row {row_number}: " + " | ".join(f"{get_column_letter(i)}={value}" for i, value in enumerate(values, start=1)))
    return ExcelText(source=path, text="\n".join(lines), sheet_count=1, row_count=row_count)


def render_excel_file(path: str | Path) -> ExcelText:
    source = Path(path)
    if source.suffix.lower() == ".csv":
        return _render_csv(source)

    workbook = load_workbook(source, data_only=True, read_only=False)
    lines = [f"# Source: {source.name}"]
    total_rows = 0
    try:
        for sheet in workbook.worksheets:
            merged_values: dict[tuple[int, int], str] = {}
            merged_notes: list[str] = []
            for merged in sheet.merged_cells.ranges:
                anchor = sheet.cell(merged.min_row, merged.min_col).value
                anchor_text = cell_text(anchor)
                if anchor_text:
                    merged_notes.append(f"{merged.coord}={anchor_text}")
                for row in range(merged.min_row, merged.max_row + 1):
                    for col in range(merged.min_col, merged.max_col + 1):
                        merged_values[(row, col)] = anchor_text

            sheet_lines = [f"## Sheet: {sheet.title}"]
            if merged_notes:
                sheet_lines.append("merged: " + " | ".join(merged_notes))

            used_rows = 0
            for row in sheet.iter_rows():
                values: list[str] = []
                for cell in row:
                    value = cell_text(cell.value)
                    if not value:
                        value = merged_values.get((cell.row, cell.column), "")
                    values.append(value)
                values = _trim_cells(values)
                if not values:
                    continue
                used_rows += 1
                total_rows += 1
                cells = " | ".join(
                    f"{get_column_letter(index)}={value}"
                    for index, value in enumerate(values, start=1)
                    if value != ""
                )
                sheet_lines.append(f"row {row[0].row}: {cells}")

            if used_rows:
                lines.extend(sheet_lines)
    finally:
        workbook.close()

    return ExcelText(source=source, text="\n".join(lines), sheet_count=len(workbook.worksheets), row_count=total_rows)
